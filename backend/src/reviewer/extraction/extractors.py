from __future__ import annotations

import csv
import importlib
import io
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook

from reviewer.extraction.models import ExtractedBlock, ExtractedDocument


class UnsupportedDocumentError(ValueError):
    pass


class Extractor(Protocol):
    extensions: frozenset[str]

    def extract(self, filename: str, content: bytes) -> ExtractedDocument: ...


def _block(
    text: str,
    source: str,
    offset: int,
    **location: int | str | None,
) -> ExtractedBlock:
    page = location.get("page")
    sheet = location.get("sheet")
    cell_range = location.get("cell_range")
    return ExtractedBlock(
        text=text,
        source=source,
        start_offset=offset,
        end_offset=offset + len(text),
        page=page if isinstance(page, int) else None,
        sheet=sheet if isinstance(sheet, str) else None,
        cell_range=cell_range if isinstance(cell_range, str) else None,
        kind=str(location.get("kind") or "text"),
    )


class _OcrHelper:
    """Optional OCR helper for images and scanned PDFs.

    Uses pytesseract + Pillow when the [ocr] optional dependency group is
    installed and the Tesseract binary is available. The helper loads lazily
    and fails silently so the rest of the app can run without OCR.
    """

    def __init__(self) -> None:
        self._available = False
        self._pytesseract: Any | None = None
        self._Image: Any | None = None
        try:
            self._pytesseract = importlib.import_module("pytesseract")
            self._Image = importlib.import_module("PIL.Image")
            self._pytesseract.get_tesseract_version()
            self._available = True
        except Exception:
            pass

    def read_image(self, image_bytes: bytes) -> str | None:
        if not self._available or self._Image is None or self._pytesseract is None:
            return None
        image = self._Image.open(io.BytesIO(image_bytes))
        return str(self._pytesseract.image_to_string(image)).strip()


_CONNECTORS = frozenset({"and", "&", "+", "or"})


def _is_number_line(text: str) -> bool:
    text = text.strip()
    if not text:
        return False
    from reviewer.analysis.financial import NUMBER_PATTERN

    return NUMBER_PATTERN.fullmatch(text) is not None


def _merge_text_lines(page_text: str) -> list[str]:
    """Merge wrapped connector lines and label-number pairs in extracted text."""
    raw_lines = [line for line in page_text.splitlines() if line.strip()]
    merged_lines: list[str] = []
    i = 0
    while i < len(raw_lines):
        line = raw_lines[i].strip()
        i += 1
        while i < len(raw_lines) and line.split()[-1].lower() in _CONNECTORS:
            line = f"{line} {raw_lines[i].strip()}"
            i += 1
        while i < len(raw_lines) and _is_number_line(raw_lines[i]):
            line = f"{line} {raw_lines[i].strip()}"
            i += 1
        merged_lines.append(line)
    return merged_lines


class PlainTextExtractor:
    extensions = frozenset({".txt", ".md", ".json"})

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        text = content.decode("utf-8-sig")
        return ExtractedDocument(
            filename=filename,
            parser="plain-text",
            blocks=[_block(text, filename, 0)],
        )


class CsvExtractor:
    extensions = frozenset({".csv"})

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        blocks: list[ExtractedBlock] = []
        offset = 0
        for row_number, row in enumerate(rows, start=1):
            text = " | ".join(value.strip() for value in row)
            blocks.append(
                _block(
                    text,
                    filename,
                    offset,
                    kind="table-row",
                    cell_range=f"{row_number}:{row_number}",
                )
            )
            offset += len(text) + 1
        return ExtractedDocument(filename=filename, parser="python-csv", blocks=blocks)


class PdfExtractor:
    extensions = frozenset({".pdf"})

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        import fitz  # type: ignore[import-untyped]

        blocks: list[ExtractedBlock] = []
        warnings: list[str] = []
        offset = 0
        ocr = _OcrHelper()

        with fitz.open(stream=content, filetype="pdf") as document:
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text")
                if not page_text.strip():
                    if ocr._available:
                        try:
                            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                            ocr_text = ocr.read_image(pixmap.tobytes("png"))
                        except Exception:
                            ocr_text = None
                        if ocr_text:
                            page_text = ocr_text
                        else:
                            warnings.append(
                                f"Page {page_number} has no embedded text and OCR produced no text."
                            )
                            continue
                    else:
                        warnings.append(
                            f"Page {page_number} has no embedded text; OCR is required."
                        )
                        continue

                merged_lines = _merge_text_lines(page_text)
                for line_no, line in enumerate(merged_lines, start=1):
                    blocks.append(
                        _block(
                            line,
                            filename,
                            offset,
                            page=page_number,
                            cell_range=f"L{line_no}",
                            kind="line",
                        )
                    )
                    offset += len(line) + 1
        return ExtractedDocument(
            filename=filename,
            parser="pymupdf",
            blocks=blocks,
            warnings=warnings,
        )


class DocxExtractor:
    extensions = frozenset({".docx"})

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        from docx import Document as WordDocument

        document = WordDocument(io.BytesIO(content))
        blocks: list[ExtractedBlock] = []
        offset = 0
        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                blocks.append(_block(text, filename, offset, cell_range=f"P{paragraph_number}"))
                offset += len(text) + 1
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                blocks.append(
                    _block(
                        text,
                        filename,
                        offset,
                        kind="table-row",
                        cell_range=f"T{table_number}:R{row_number}",
                    )
                )
                offset += len(text) + 1
        return ExtractedDocument(filename=filename, parser="python-docx", blocks=blocks)


class XlsxExtractor:
    extensions = frozenset({".xlsx"})

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        blocks: list[ExtractedBlock] = []
        offset = 0
        for sheet in workbook.worksheets:
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                rendered = ["" if value is None else str(value) for value in values]
                if not any(rendered):
                    continue
                text = " | ".join(rendered)
                blocks.append(
                    _block(
                        text,
                        filename,
                        offset,
                        kind="table-row",
                        sheet=sheet.title,
                        cell_range=f"{row_number}:{row_number}",
                    )
                )
                offset += len(text) + 1
        return ExtractedDocument(filename=filename, parser="openpyxl", blocks=blocks)


class DoclingExtractor:
    extensions = frozenset({".pdf"})

    def __init__(self) -> None:
        import importlib

        self._converter: Any | None = None
        try:
            docling_mod = importlib.import_module("docling.document_converter")
            converter = docling_mod.DocumentConverter
        except Exception:
            self.extensions = frozenset()
        else:
            self._converter = converter

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        import os
        import tempfile

        if self._converter is None:
            raise UnsupportedDocumentError("Docling is not installed.")

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name

        try:
            result = self._converter().convert(tmp_path)
            doc = result.document
            page_contents: dict[int, list[str]] = {}
            for item in doc.texts:
                if not item.text:
                    continue
                page = item.prov[0].page if item.prov else 1
                page_contents.setdefault(page, []).append(item.text)
            for table in doc.tables:
                table_md = table.export_to_markdown(doc)
                if not table_md:
                    continue
                page = table.prov[0].page if table.prov else 1
                page_contents.setdefault(page, []).append(table_md)

            blocks: list[ExtractedBlock] = []
            offset = 0
            for page in sorted(page_contents):
                page_text = "\n".join(page_contents[page])
                blocks.append(
                    _block(
                        page_text,
                        filename,
                        offset,
                        page=page,
                        kind="docling-page",
                    )
                )
                offset += len(page_text) + 1
            return ExtractedDocument(
                filename=filename,
                parser="docling",
                blocks=blocks,
                warnings=[],
            )
        except Exception as exc:
            raise UnsupportedDocumentError(f"Docling conversion failed: {exc}") from exc
        finally:
            os.unlink(tmp_path)


class ImageExtractor:
    extensions = frozenset({".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".gif"})

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        ocr = _OcrHelper()
        if not ocr._available:
            raise UnsupportedDocumentError(
                "OCR is not installed. Add the [ocr] optional dependency group."
            )

        text = ocr.read_image(content)
        warnings: list[str] = []
        if not text:
            warnings.append("No text was detected in the image.")

        blocks: list[ExtractedBlock] = []
        offset = 0
        if text:
            merged_lines = _merge_text_lines(text)
            for line_no, line in enumerate(merged_lines, start=1):
                blocks.append(
                    _block(
                        line,
                        filename,
                        offset,
                        page=1,
                        cell_range=f"L{line_no}",
                        kind="ocr-line",
                    )
                )
                offset += len(line) + 1

        return ExtractedDocument(
            filename=filename,
            parser="ocr",
            blocks=blocks,
            warnings=warnings,
        )


class ExtractorRegistry:
    def __init__(self, extractors: list[Extractor] | None = None) -> None:
        self._extractors = extractors or [
            PlainTextExtractor(),
            CsvExtractor(),
            DoclingExtractor(),
            PdfExtractor(),
            ImageExtractor(),
            DocxExtractor(),
            XlsxExtractor(),
        ]

    @property
    def supported_extensions(self) -> set[str]:
        return {extension for extractor in self._extractors for extension in extractor.extensions}

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        extension = Path(filename).suffix.lower()
        for extractor in self._extractors:
            if extension in extractor.extensions:
                try:
                    return extractor.extract(filename, content)
                except UnsupportedDocumentError:
                    continue
        supported = ", ".join(sorted(self.supported_extensions))
        raise UnsupportedDocumentError(
            f"Unsupported file type {extension!r}; expected {supported}."
        )
